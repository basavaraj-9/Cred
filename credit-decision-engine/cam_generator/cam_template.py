from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors
from datetime import datetime
import os
from typing import Dict, Any

class CAMGenerator:
    def __init__(self):
        self.template_path = "cam_generator/templates/"
        self.output_path = "generated_cams/"
        
    def generate_cam(self, company_data: Dict, financial_data: Dict, risk_analysis: Dict,
                     research_data: Dict, recommendation: Dict) -> Dict[str, str]:
        """Generate Credit Appraisal Memo in both Word and PDF formats"""
        
        # Ensure output directory exists
        os.makedirs(self.output_path, exist_ok=True)
        
        # Generate filename
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        company_name = company_data.get('company', 'Unknown').replace(' ', '_')
        base_filename = f"CAM_{company_name}_{timestamp}"
        
        # Generate Word document
        word_path = os.path.join(self.output_path, f"{base_filename}.docx")
        self._generate_word_cam(word_path, company_data, financial_data, risk_analysis, 
                               research_data, recommendation)
        
        # Generate PDF document
        pdf_path = os.path.join(self.output_path, f"{base_filename}.pdf")
        self._generate_pdf_cam(pdf_path, company_data, financial_data, risk_analysis, 
                              research_data, recommendation)
        
        return {
            'word_path': word_path,
            'pdf_path': pdf_path,
            'filename': base_filename
        }
    
    def _generate_word_cam(self, filepath: str, company_data: Dict, financial_data: Dict,
                          risk_analysis: Dict, research_data: Dict, recommendation: Dict):
        """Generate CAM in Word format"""
        
        doc = Document()
        
        # Title
        title = doc.add_heading('CREDIT APPRAISAL MEMORANDUM', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Date and Reference
        doc.add_paragraph(f'Date: {datetime.now().strftime("%d %B %Y")}')
        doc.add_paragraph(f'Reference: CAM/{company_data.get("company", "XXX")}/{datetime.now().strftime("%Y%m%d")}')
        doc.add_paragraph('')
        
        # 1. Borrower Overview
        doc.add_heading('1. BORROWER OVERVIEW', level=1)
        self._add_borrower_overview(doc, company_data, financial_data)
        
        # 2. Industry Analysis
        doc.add_heading('2. INDUSTRY ANALYSIS', level=1)
        self._add_industry_analysis(doc, research_data)
        
        # 3. Financial Analysis
        doc.add_heading('3. FINANCIAL ANALYSIS', level=1)
        self._add_financial_analysis(doc, financial_data, risk_analysis)
        
        # 4. Risk Assessment
        doc.add_heading('4. RISK ASSESSMENT', level=1)
        self._add_risk_assessment(doc, risk_analysis, research_data)
        
        # 5. Five Cs of Credit
        doc.add_heading('5. FIVE CS OF CREDIT', level=1)
        self._add_five_cs(doc, company_data, financial_data, risk_analysis, research_data)
        
        # 6. Recommendation
        doc.add_heading('6. RECOMMENDATION', level=1)
        self._add_recommendation(doc, recommendation)
        
        # Save document
        doc.save(filepath)
    
    def _generate_pdf_cam(self, filepath: str, company_data: Dict, financial_data: Dict,
                         risk_analysis: Dict, research_data: Dict, recommendation: Dict):
        """Generate CAM in PDF format"""
        
        doc = SimpleDocTemplate(filepath, pagesize=A4)
        styles = getSampleStyleSheet()
        story = []
        
        # Custom styles
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            spaceAfter=30,
            alignment=1  # Center alignment
        )
        
        # Title
        story.append(Paragraph("CREDIT APPRAISAL MEMORANDUM", title_style))
        story.append(Spacer(1, 20))
        
        # Date and Reference
        story.append(Paragraph(f"Date: {datetime.now().strftime('%d %B %Y')}", styles['Normal']))
        story.append(Paragraph(f"Reference: CAM/{company_data.get('company', 'XXX')}/{datetime.now().strftime('%Y%m%d')}", styles['Normal']))
        story.append(Spacer(1, 20))
        
        # 1. Borrower Overview
        story.append(Paragraph("1. BORROWER OVERVIEW", styles['Heading1']))
        story.extend(self._get_borrower_overview_pdf(company_data, financial_data, styles))
        
        # 2. Industry Analysis
        story.append(Paragraph("2. INDUSTRY ANALYSIS", styles['Heading1']))
        story.extend(self._get_industry_analysis_pdf(research_data, styles))
        
        # 3. Financial Analysis
        story.append(Paragraph("3. FINANCIAL ANALYSIS", styles['Heading1']))
        story.extend(self._get_financial_analysis_pdf(financial_data, risk_analysis, styles))
        
        # 4. Risk Assessment
        story.append(Paragraph("4. RISK ASSESSMENT", styles['Heading1']))
        story.extend(self._get_risk_assessment_pdf(risk_analysis, research_data, styles))
        
        # 5. Five Cs of Credit
        story.append(Paragraph("5. FIVE CS OF CREDIT", styles['Heading1']))
        story.extend(self._get_five_cs_pdf(company_data, financial_data, risk_analysis, research_data, styles))
        
        # 6. Recommendation
        story.append(Paragraph("6. RECOMMENDATION", styles['Heading1']))
        story.extend(self._get_recommendation_pdf(recommendation, styles))
        
        # Build PDF
        doc.build(story)
    
    def _add_borrower_overview(self, doc, company_data: Dict, financial_data: Dict):
        """Add borrower overview section"""
        
        p = doc.add_paragraph()
        p.add_run('Company Name: ').bold = True
        p.add_run(company_data.get('company', 'N/A'))
        
        p = doc.add_paragraph()
        p.add_run('Business Nature: ').bold = True
        p.add_run(company_data.get('business_nature', 'Manufacturing/Trading/Services'))
        
        p = doc.add_paragraph()
        p.add_run('Years in Operation: ').bold = True
        p.add_run(str(company_data.get('years_in_operation', 'N/A')))
        
        p = doc.add_paragraph()
        p.add_run('Annual Revenue: ').bold = True
        p.add_run(f"₹{financial_data.get('revenue', 0):,}")
        
        p = doc.add_paragraph()
        p.add_run('Existing Loans: ').bold = True
        p.add_run(f"₹{financial_data.get('existing_loans', 0):,}")
        
        doc.add_paragraph('')
    
    def _add_industry_analysis(self, doc, research_data: Dict):
        """Add industry analysis section"""
        
        p = doc.add_paragraph()
        p.add_run('Sector: ').bold = True
        p.add_run(research_data.get('sector', 'N/A'))
        
        p = doc.add_paragraph()
        p.add_run('Sector Outlook: ').bold = True
        p.add_run(research_data.get('sector_outlook', 'Neutral'))
        
        p = doc.add_paragraph()
        p.add_run('Market Sentiment: ').bold = True
        p.add_run(research_data.get('news_sentiment', 'Neutral'))
        
        p = doc.add_paragraph()
        p.add_run('Key Industry Trends: ').bold = True
        doc.add_paragraph('• Digital transformation and automation')
        doc.add_paragraph('• Regulatory compliance requirements')
        doc.add_paragraph('• Supply chain disruptions and mitigation')
        
        doc.add_paragraph('')
    
    def _add_financial_analysis(self, doc, financial_data: Dict, risk_analysis: Dict):
        """Add financial analysis section"""
        
        # Financial metrics table
        table_data = [
            ['Financial Metric', 'Amount (₹)', 'Analysis'],
            ['Revenue', f"{financial_data.get('revenue', 0):,}", self._analyze_revenue(financial_data)],
            ['Total Assets', f"{financial_data.get('assets', 0):,}", self._analyze_assets(financial_data)],
            ['Total Liabilities', f"{financial_data.get('liabilities', 0):,}", self._analyze_liabilities(financial_data)],
            ['Existing Loans', f"{financial_data.get('existing_loans', 0):,}", self._analyze_loans(financial_data)],
            ['Net Profit', f"{financial_data.get('profit', 0):,}", self._analyze_profit(financial_data)]
        ]
        
        table = doc.add_table(rows=len(table_data), cols=3)
        table.style = 'Table Grid'
        
        for i, row in enumerate(table_data):
            for j, cell_text in enumerate(row):
                table.cell(i, j).text = cell_text
                if i == 0:  # Header row
                    table.cell(i, j).paragraphs[0].runs[0].bold = True
        
        doc.add_paragraph('')
        
        # Financial ratios
        p = doc.add_paragraph()
        p.add_run('Key Financial Ratios: ').bold = True
        
        ratios = self._calculate_financial_ratios(financial_data)
        for ratio_name, ratio_value in ratios.items():
            doc.add_paragraph(f"• {ratio_name}: {ratio_value}")
        
        doc.add_paragraph('')
    
    def _add_risk_assessment(self, doc, risk_analysis: Dict, research_data: Dict):
        """Add risk assessment section"""
        
        p = doc.add_paragraph()
        p.add_run('Overall Risk Score: ').bold = True
        p.add_run(f"{risk_analysis.get('risk_score', 0)}/100")
        
        p = doc.add_paragraph()
        p.add_run('Risk Category: ').bold = True
        p.add_run(risk_analysis.get('risk_category', 'Medium'))
        
        p = doc.add_paragraph()
        p.add_run('Risk Factors Identified: ').bold = True
        
        risk_factors = risk_analysis.get('risk_factors', [])
        for factor in risk_factors:
            doc.add_paragraph(f"• {factor}")
        
        if not risk_factors:
            doc.add_paragraph("• No significant risk factors identified")
        
        doc.add_paragraph('')
        
        # Component scores
        p = doc.add_paragraph()
        p.add_run('Component Risk Scores: ').bold = True
        
        component_scores = risk_analysis.get('component_scores', {})
        for component, score in component_scores.items():
            doc.add_paragraph(f"• {component.title()}: {score}/100")
        
        doc.add_paragraph('')
    
    def _add_five_cs(self, doc, company_data: Dict, financial_data: Dict,
                     risk_analysis: Dict, research_data: Dict):
        """Add Five Cs of Credit analysis"""
        
        # Character
        doc.add_heading('Character', level=2)
        character_score = self._assess_character(company_data, research_data)
        doc.add_paragraph(f"Assessment: {character_score}")
        doc.add_paragraph("")
        
        # Capacity
        doc.add_heading('Capacity', level=2)
        capacity_score = self._assess_capacity(financial_data, risk_analysis)
        doc.add_paragraph(f"Assessment: {capacity_score}")
        doc.add_paragraph("")
        
        # Capital
        doc.add_heading('Capital', level=2)
        capital_score = self._assess_capital(financial_data)
        doc.add_paragraph(f"Assessment: {capital_score}")
        doc.add_paragraph("")
        
        # Collateral
        doc.add_heading('Collateral', level=2)
        collateral_score = self._assess_collateral(financial_data, risk_analysis)
        doc.add_paragraph(f"Assessment: {collateral_score}")
        doc.add_paragraph("")
        
        # Conditions
        doc.add_heading('Conditions', level=2)
        conditions_score = self._assess_conditions(research_data, risk_analysis)
        doc.add_paragraph(f"Assessment: {conditions_score}")
        doc.add_paragraph("")
    
    def _add_recommendation(self, doc, recommendation: Dict):
        """Add recommendation section"""
        
        p = doc.add_paragraph()
        p.add_run('Decision: ').bold = True
        p.add_run(recommendation.get('decision', 'N/A'))
        
        p = doc.add_paragraph()
        p.add_run('Recommended Loan Amount: ').bold = True
        p.add_run(f"₹{recommendation.get('loan_limit', 0):,}")
        
        p = doc.add_paragraph()
        p.add_run('Interest Rate: ').bold = True
        p.add_run(f"{recommendation.get('interest_rate', 0)}% p.a.")
        
        p = doc.add_paragraph()
        p.add_run('Tenure: ').bold = True
        p.add_run(f"{recommendation.get('tenure_months', 0)} months")
        
        p = doc.add_paragraph()
        p.add_run('Conditions: ').bold = True
        
        conditions = recommendation.get('conditions', [])
        for condition in conditions:
            doc.add_paragraph(f"• {condition}")
        
        doc.add_paragraph('')
        
        p = doc.add_paragraph()
        p.add_run('Justification: ').bold = True
        
        explanations = recommendation.get('explanations', [])
        for explanation in explanations:
            doc.add_paragraph(f"• {explanation}")
    
    # Helper methods for PDF generation
    def _get_borrower_overview_pdf(self, company_data: Dict, financial_data: Dict, styles):
        """Get borrower overview content for PDF"""
        content = []
        
        data = [
            ['Company Name:', company_data.get('company', 'N/A')],
            ['Business Nature:', company_data.get('business_nature', 'Manufacturing/Trading/Services')],
            ['Annual Revenue:', f"₹{financial_data.get('revenue', 0):,}"],
            ['Existing Loans:', f"₹{financial_data.get('existing_loans', 0):,}"]
        ]
        
        table = Table(data, colWidths=[2*inch, 4*inch])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.lightgrey),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
        ]))
        
        content.append(table)
        content.append(Spacer(1, 20))
        
        return content
    
    def _get_industry_analysis_pdf(self, research_data: Dict, styles):
        """Get industry analysis content for PDF"""
        content = []
        
        sector = research_data.get('sector', 'N/A')
        outlook = research_data.get('sector_outlook', 'Neutral')
        
        content.append(Paragraph(f"<b>Sector:</b> {sector}", styles['Normal']))
        content.append(Paragraph(f"<b>Sector Outlook:</b> {outlook}", styles['Normal']))
        content.append(Spacer(1, 12))
        
        return content
    
    def _get_financial_analysis_pdf(self, financial_data: Dict, risk_analysis: Dict, styles):
        """Get financial analysis content for PDF"""
        content = []
        
        data = [['Metric', 'Amount (₹)', 'Analysis']]
        data.append([
            'Revenue',
            f"{financial_data.get('revenue', 0):,}",
            self._analyze_revenue(financial_data)
        ])
        data.append([
            'Total Assets',
            f"{financial_data.get('assets', 0):,}",
            self._analyze_assets(financial_data)
        ])
        
        table = Table(data, colWidths=[2*inch, 2*inch, 2*inch])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        
        content.append(table)
        content.append(Spacer(1, 20))
        
        return content
    
    def _get_risk_assessment_pdf(self, risk_analysis: Dict, research_data: Dict, styles):
        """Get risk assessment content for PDF"""
        content = []
        
        content.append(Paragraph(f"<b>Risk Score:</b> {risk_analysis.get('risk_score', 0)}/100", styles['Normal']))
        content.append(Paragraph(f"<b>Risk Category:</b> {risk_analysis.get('risk_category', 'Medium')}", styles['Normal']))
        content.append(Spacer(1, 12))
        
        return content
    
    def _get_five_cs_pdf(self, company_data: Dict, financial_data: Dict, 
                        risk_analysis: Dict, research_data: Dict, styles):
        """Get Five Cs content for PDF"""
        content = []
        
        content.append(Paragraph("<b>Character:</b> " + self._assess_character(company_data, research_data), styles['Normal']))
        content.append(Paragraph("<b>Capacity:</b> " + self._assess_capacity(financial_data, risk_analysis), styles['Normal']))
        content.append(Paragraph("<b>Capital:</b> " + self._assess_capital(financial_data), styles['Normal']))
        content.append(Spacer(1, 12))
        
        return content
    
    def _get_recommendation_pdf(self, recommendation: Dict, styles):
        """Get recommendation content for PDF"""
        content = []
        
        content.append(Paragraph(f"<b>Decision:</b> {recommendation.get('decision', 'N/A')}", styles['Normal']))
        content.append(Paragraph(f"<b>Loan Amount:</b> ₹{recommendation.get('loan_limit', 0):,}", styles['Normal']))
        content.append(Paragraph(f"<b>Interest Rate:</b> {recommendation.get('interest_rate', 0)}% p.a.", styles['Normal']))
        content.append(Spacer(1, 12))
        
        return content
    
    # Helper methods for analysis
    def _analyze_revenue(self, financial_data: Dict) -> str:
        """Analyze revenue"""
        revenue = financial_data.get('revenue', 0)
        if revenue > 100000000:  # > 10 crore
            return "Strong revenue base"
        elif revenue > 10000000:  # > 1 crore
            return "Adequate revenue"
        else:
            return "Limited revenue"
    
    def _analyze_assets(self, financial_data: Dict) -> str:
        """Analyze assets"""
        assets = financial_data.get('assets', 0)
        liabilities = financial_data.get('liabilities', 0)
        
        if assets > 0 and liabilities > 0:
            ratio = liabilities / assets
            if ratio < 0.5:
                return "Healthy asset base"
            elif ratio < 0.8:
                return "Moderate asset utilization"
            else:
                return "High leverage on assets"
        return "Insufficient asset data"
    
    def _analyze_liabilities(self, financial_data: Dict) -> str:
        """Analyze liabilities"""
        liabilities = financial_data.get('liabilities', 0)
        revenue = financial_data.get('revenue', 0)
        
        if revenue > 0:
            ratio = liabilities / revenue
            if ratio < 0.5:
                return "Manageable liability level"
            elif ratio < 1.0:
                return "Moderate liability burden"
            else:
                return "High liability burden"
        return "Cannot assess liability ratio"
    
    def _analyze_loans(self, financial_data: Dict) -> str:
        """Analyze existing loans"""
        loans = financial_data.get('existing_loans', 0)
        revenue = financial_data.get('revenue', 0)
        
        if revenue > 0:
            ratio = loans / revenue
            if ratio < 0.3:
                return "Low existing debt"
            elif ratio < 0.7:
                return "Moderate existing debt"
            else:
                return "High existing debt burden"
        return "Cannot assess debt ratio"
    
    def _analyze_profit(self, financial_data: Dict) -> str:
        """Analyze profit"""
        profit = financial_data.get('profit', 0)
        revenue = financial_data.get('revenue', 0)
        
        if revenue > 0:
            margin = profit / revenue
            if margin > 0.15:
                return "Strong profitability"
            elif margin > 0.05:
                return "Adequate profitability"
            elif margin > 0:
                return "Low profitability"
            else:
                return "Negative profitability"
        return "Cannot assess profit margin"
    
    def _calculate_financial_ratios(self, financial_data: Dict) -> Dict[str, str]:
        """Calculate key financial ratios"""
        ratios = {}
        
        revenue = financial_data.get('revenue', 0)
        assets = financial_data.get('assets', 0)
        liabilities = financial_data.get('liabilities', 0)
        profit = financial_data.get('profit', 0)
        existing_loans = financial_data.get('existing_loans', 0)
        
        if revenue > 0:
            ratios['Debt to Revenue'] = f"{existing_loans/revenue:.2f}x"
            ratios['Profit Margin'] = f"{(profit/revenue)*100:.1f}%"
        
        if assets > 0:
            ratios['Debt to Assets'] = f"{liabilities/assets:.2f}x"
        
        return ratios
    
    def _assess_character(self, company_data: Dict, research_data: Dict) -> str:
        """Assess character"""
        promoter_risk = research_data.get('promoter_risk', 'Medium')
        litigation_cases = research_data.get('litigation_cases', 0)
        
        if promoter_risk == 'Low' and litigation_cases == 0:
            return "Strong character and reputation"
        elif promoter_risk == 'Medium' and litigation_cases <= 2:
            return "Acceptable character with minor concerns"
        else:
            return "Character concerns requiring attention"
    
    def _assess_capacity(self, financial_data: Dict, risk_analysis: Dict) -> str:
        """Assess capacity to repay"""
        revenue = financial_data.get('revenue', 0)
        profit = financial_data.get('profit', 0)
        
        if profit > 0 and revenue > 10000000:
            return "Strong repayment capacity"
        elif profit >= 0:
            return "Adequate repayment capacity"
        else:
            return "Limited repayment capacity"
    
    def _assess_capital(self, financial_data: Dict) -> str:
        """Assess capital base"""
        assets = financial_data.get('assets', 0)
        liabilities = financial_data.get('liabilities', 0)
        
        if assets > 0 and liabilities < assets * 0.5:
            return "Strong capital base"
        elif assets > 0:
            return "Adequate capital base"
        else:
            return "Limited capital base"
    
    def _assess_collateral(self, financial_data: Dict, risk_analysis: Dict) -> str:
        """Assess collateral availability"""
        assets = financial_data.get('assets', 0)
        risk_category = risk_analysis.get('risk_category', 'Medium')
        
        if assets > 50000000:  # > 5 crore
            return "Sufficient collateral available"
        elif assets > 10000000:  # > 1 crore
            return "Moderate collateral available"
        else:
            return "Limited collateral"
    
    def _assess_conditions(self, research_data: Dict, risk_analysis: Dict) -> str:
        """Assess economic and industry conditions"""
        sector_outlook = research_data.get('sector_outlook', 'Neutral')
        
        if sector_outlook == 'Positive':
            return "Favorable economic conditions"
        elif sector_outlook == 'Negative':
            return "Challenging economic conditions"
        else:
            return "Stable economic conditions"
